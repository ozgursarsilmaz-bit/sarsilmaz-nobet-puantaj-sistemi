import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- 1. MEVCUT EKİM 2026 NÖBET DİZİLİMİ ---
october_schedule = [
    ("01.10.2026", "Perşembe", "ŞADUMAN YALÇIN (Acil)", "DURMUŞ AKTÜRK", "M.ATABERK YILDIZ"),
    ("02.10.2026", "Cuma", "MURAT GENCER", "ÖZGÜR SARSILMAZ", "GÜLÇİN HORDACI (Acil)"),
    ("03.10.2026", "Cumartesi", "ŞADUMAN YALÇIN", "KADRİYE ARDIÇ (Acil)", "HİCRAN ATA AYHAN"),
    ("04.10.2026", "Pazar", "SUNA SARSILMAZ", "ŞENGÜL URLU (Acil)", "HACER BÜKÜM"),
    ("05.10.2026", "Pazartesi", "ŞADUMAN YALÇIN", "HÜMEYRA YEŞİLTAŞ (Acil)", "BİRSEL KAYA"),
    ("06.10.2026", "Salı", "MURAT GENCER", "REMZİYE AKSOY (Acil)", "İBRAHİM ER"),
    ("07.10.2026", "Çarşamba", "ŞADUMAN YALÇIN", "DURMUŞ AKTÜRK (Acil)", "M.ATABERK YILDIZ"),
    ("08.10.2026", "Perşembe", "SUNA SARSILMAZ (Acil)", "MELEK AKKAŞ", "GÜLÇİN HORDACI"),
    ("09.10.2026", "Cuma", "KEVSER DUMLU (Acil)", "ŞENGÜL URLU", "HACER BÜKÜM"),
    ("10.10.2026", "Cumartesi", "SEÇİL EMİR YILDIRAK", "MÜCELLÂ MANTAR", "BİRSEL KAYA (Acil)"),
    ("11.10.2026", "Pazar", "BELGİN UYSAL", "ÖZTÜRK MAVİŞ", "HİCRAN ATA AYHAN (Acil)"),
    ("12.10.2026", "Pazartesi", "MURAT GENCER", "REMZİYE AKSOY", "SÜLEYMAN POLAT (Acil)"),
    ("13.10.2026", "Salı", "ŞENEL TAŞ (Acil)", "KADRİYE ARDIÇ", "GÜLÇİN HORDACI"),
    ("14.10.2026", "Çarşamba", "KEVSER DUMLU", "MELEK AKKAŞ", "HACER BÜKÜM (Acil)"),
    ("15.10.2026", "Perşembe", "BELGİN UYSAL", "ŞÜKRAN KAYA", "MELEK BARUT (Acil)"),
    ("16.10.2026", "Cuma", "SUNA SARSILMAZ", "ÖZTÜRK MAVİŞ", "M.ATABERK YILDIZ (Acil)"),
    ("17.10.2026", "Cumartesi", "ŞENEL TAŞ", "ÖZGÜR SARSILMAZ (Acil)", "SÜLEYMAN POLAT"),
    ("18.10.2026", "Pazar", "MURAT GENCER", "REMZİYE AKSOY", "İBRAHİM ER (Acil)"),
    ("19.10.2026", "Pazartesi", "KEVSER DUMLU", "MUSTAFA TOSUN (Acil)", "GÜLÇİN HORDACI"),
    ("20.10.2026", "Salı", "SEÇİL EMİR YILDIRAK", "DURMUŞ AKTÜRK", "HİCRAN ATA AYHAN (Acil)"),
    ("21.10.2026", "Çarşamba", "ŞENEL TAŞ", "ÖZTÜRK MAVİŞ (Acil)", "MELEK BARUT"),
    ("22.10.2026", "Perşembe", "MURAT GENCER (Acil)", "MELEK AKKAŞ", "HACER BÜKÜM"),
    ("23.10.2026", "Cuma", "BELGİN UYSAL (Acil)", "HÜMEYRA YEŞİLTAŞ", "İBRAHİM ER"),
    ("24.10.2026", "Cumartesi", "SUNA SARSILMAZ", "MUSTAFA TOSUN", "MUSTAFA HARTOĞLU (Acil)"),
    ("25.10.2026", "Pazar", "ŞENEL TAŞ", "MEHMET BAĞCI (Acil)", "SÜLEYMAN POLAT"),
    ("26.10.2026", "Pazartesi", "SEÇİL EMİR YILDIRAK", "DURMUŞ AKTÜRK (Acil)", "BİRSEL KAYA"),
    ("27.10.2026", "Salı", "KEVSER DUMLU", "MELEK AKKAŞ (Acil)", "M.ATABERK YILDIZ"),
    ("28.10.2026", "Çarşamba", "BELGİN UYSAL", "MÜCELLÂ MANTAR (Acil)", "GÜLÇİN HORDACI"),
    ("29.10.2026", "Perşembe", "ŞENEL TAŞ", "HÜMEYRA YEŞİLTAŞ (Acil)", "İBRAHİM ER"),
    ("30.10.2026", "Cuma", "SEÇİL EMİR YILDIRAK (Acil)", "MEHMET BAĞCI", "MUSTAFA HARTOĞLU"),
    ("31.10.2026", "Cumartesi", "KEVSER DUMLU", "ŞÜKRAN KAYA (Acil)", "MELEK BARUT")
]

TUM_PERSONELLER = [
    "BELGİN UYSAL", "BİRSEL KAYA", "DURMUŞ AKTÜRK", "GÜLÇİN HORDACI", "HACER BÜKÜM",
    "HİCRAN ATA AYHAN", "HÜMEYRA YEŞİLTAŞ", "İBRAHİM ER", "KADRİYE ARDIÇ", "KEVSER DUMLU",
    "MEHMET BAĞCI", "MELEK AKKAŞ", "MELEK BARUT", "MUHAMMED ATABERK YILDIZ", "MURAT GENCER",
    "MUSTAFA HARTOĞLU", "MUSTAFA TOSUN", "MÜCELLÂ MANTAR", "ÖZGÜR SARSILMAZ", "ÖZTÜRK MAVİŞ",
    "REMZİYE AKSOY", "SEÇİL EMİR YILDIRAK", "SEMRA KUMRUOĞLU", "SUNA SARSILMAZ", "SÜLEYMAN POLAT",
    "ŞADUMAN YALÇIN", "ŞENEL TAŞ", "ŞENGÜL URLU", "ŞÜKRAN KAYA"
]

gunler_listesi = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

# Nöbet Ayrıştırma Haritası
nobet_map = {p: {} for p in TUM_PERSONELLER}
for idx, (tarih_str, gun_str, pcr, mikro, kultur) in enumerate(october_schedule, 1):
    for p_entry in [pcr, mikro, kultur]:
        is_acil = "(Acil)" in p_entry
        p_name = p_entry.replace(" (Acil)", "").strip().upper()
        if p_name == "M.ATABERK YILDIZ":
            p_name = "MUHAMMED ATABERK YILDIZ"
        if p_name in nobet_map:
            nobet_map[p_name][idx] = "24A" if is_acil else "24"

wb = openpyxl.Workbook()

# Stil Tanımları
font_title = Font(name="Calibri", size=12, bold=True)
font_header = Font(name="Calibri", size=10, bold=True)
font_body = Font(name="Calibri", size=9)
font_bold = Font(name="Calibri", size=9, bold=True)
font_ni = Font(name="Calibri", size=9, bold=True, color="C00000")
font_24 = Font(name="Calibri", size=9, bold=True, color="002060")

fill_header = PatternFill(start_color="C5D9A4", end_color="C5D9A4", fill_type="solid")
fill_green = PatternFill(start_color="D8E4BC", end_color="D8E4BC", fill_type="solid")
fill_grey = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
fill_yellow = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

thin = Side(style="thin", color="A6A6A6")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_v90 = Alignment(horizontal="center", vertical="center", text_rotation=90)

# --- 1. SEKME: AYLIK GÖREV LİSTESİ ---
ws1 = wb.active
ws1.title = "Aylık Görev Listesi"
ws1.views.sheetView[0].showGridLines = True
ws1.cell(row=1, column=1, value="2026 yılı 10. Ay Mikrobiyoloji Laboratuvarı Nöbet Çizelgesi").font = font_title

headers1 = ["Tarih", "Gün", "PCR", "Mikro", "Kültür"]
for c_idx, h in enumerate(headers1, 1):
    cell = ws1.cell(row=3, column=c_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border

for idx, row in enumerate(october_schedule, 4):
    for col_i, val in enumerate(row, 1):
        c = ws1.cell(row=idx, column=col_i, value=val)
        c.font = font_body
        c.border = border
        c.alignment = align_center if col_i in [1, 2] else align_left
        if row[1] in ["Cumartesi", "Pazar"] or row[0] == "29.10.2026":
            c.fill = fill_grey

for col_l, w in zip(["A", "B", "C", "D", "E"], [13, 13, 28, 28, 28]):
    ws1.column_dimensions[col_l].width = w

# --- 2. SEKME: İSTATİSTİK & MESAİ YÜKÜ ---
ws2 = wb.create_sheet("İstatistik & Mesai Yükü")
ws2.views.sheetView[0].showGridLines = True

ws2.merge_cells("A1:A2")
ws2.cell(row=1, column=1, value="AD SOYAD").font = font_header
ws2.cell(row=1, column=1).fill = fill_header
ws2.cell(row=1, column=1).alignment = align_center
ws2.cell(row=1, column=1).border = border

col_c = 2
for g in gunler_listesi:
    ws2.merge_cells(start_row=1, start_column=col_c, end_row=1, end_column=col_c + 2)
    t_cell = ws2.cell(row=1, column=col_c, value=g)
    t_cell.font = font_header
    t_cell.fill = fill_header
    t_cell.alignment = align_center

    for sub_i, sub in enumerate(["Devir", "Bu Ay", "Toplam"]):
        sc = ws2.cell(row=2, column=col_c + sub_i, value=sub)
        sc.font = font_header
        sc.fill = fill_green
        sc.alignment = align_center
        sc.border = border
    col_c += 3

ws2.merge_cells("AB1:AB2")
c_tn = ws2.cell(row=1, column=28, value="T.NöbetSaati")
c_tn.font = font_header
c_tn.fill = fill_header
c_tn.alignment = align_center
c_tn.border = border

ws2.merge_cells("AC1:AC2")
c_tnb = ws2.cell(row=1, column=29, value="TOPLAM NÖBET")
c_tnb.font = font_header
c_tnb.fill = fill_header
c_tnb.alignment = align_center
c_tnb.border = border

for r_i in [1, 2]:
    for c_i in range(1, 30):
        ws2.cell(row=r_i, column=c_i).border = border

def calculate_shift_hours_dynamic(d):
    dt = datetime.date(2026, 10, d)
    w = dt.weekday()
    if d == 28: return 19
    if d == 27: return 11
    if d == 29: return 16
    if d in [3, 4, 10, 11, 17, 18, 24, 25, 31]: return 24
    if w in [0, 1, 2, 3]: return 16
    return 24

for p_idx, p in enumerate(TUM_PERSONELLER, 3):
    ws2.cell(row=p_idx, column=1, value=p).alignment = align_left
    ws2.cell(row=p_idx, column=1).font = font_body
    ws2.cell(row=p_idx, column=1).border = border

    bu_ay_gunler = {g: 0 for g in gunler_listesi}
    shifts = nobet_map[p]
    bu_ay_saat = 0

    for d, type_code in shifts.items():
        w_idx = datetime.date(2026, 10, d).weekday()
        bu_ay_gunler[gunler_listesi[w_idx]] += 1
        bu_ay_saat += calculate_shift_hours_dynamic(d)

    ci = 2
    for g in gunler_listesi:
        bu_ay = bu_ay_gunler[g]
        for v in [0, bu_ay, bu_ay]:
            cell = ws2.cell(row=p_idx, column=ci, value=v)
            cell.font = font_body
            cell.alignment = align_center
            cell.border = border
            ci += 1

    cs = ws2.cell(row=p_idx, column=28, value=bu_ay_saat)
    cs.font = font_bold
    cs.alignment = align_center
    cs.border = border

    cn = ws2.cell(row=p_idx, column=29, value=len(shifts))
    cn.font = font_bold
    cn.alignment = align_center
    cn.border = border

ws2.column_dimensions["A"].width = 25

# --- 3. SEKME: PUANTAJ TABLOSU ---
ws3 = wb.create_sheet("Puantaj Tablosu")
ws3.views.sheetView[0].showGridLines = True

ws3.row_dimensions[5].height = 65
ws3.cell(row=5, column=1, value="Adı Soyadı").fill = fill_green
ws3.cell(row=5, column=1).font = font_header
ws3.cell(row=5, column=1).alignment = align_left
ws3.cell(row=5, column=1).border = border

ws3.cell(row=5, column=2, value="Birim").fill = fill_green
ws3.cell(row=5, column=2).font = font_header
ws3.cell(row=5, column=2).alignment = align_center
ws3.cell(row=5, column=2).border = border

for d in range(1, 32):
    c_idx = d + 2
    dt = datetime.date(2026, 10, d)
    is_off = dt.weekday() in [5, 6] or d == 29
    cell = ws3.cell(row=5, column=c_idx, value=d)
    cell.font = font_header
    cell.alignment = align_v90
    cell.border = border
    cell.fill = fill_grey if is_off else fill_header

extra_cols = [
    "Toplam Çalışma Saati", "Aylık Çalışma Saati", "Fazla Nöbet Saati",
    "Normal Nöbet Artırımlı (Gece)", "Normal Nöbet Artırımsız (Normal)",
    "Riskli Nöbet Artırımlı (Gece)", "Riskli Nöbet Artırımsız (Normal)"
]

for col_idx, col_name in enumerate(extra_cols, 34):
    cell = ws3.cell(row=5, column=col_idx, value=col_name)
    cell.font = font_header
    cell.alignment = align_v90
    cell.border = border
    cell.fill = fill_header

for idx, p in enumerate(TUM_PERSONELLER):
    r = 6 + idx
    ws3.row_dimensions[r].height = 18

    is_muaf = (p == "SEMRA KUMRUOĞLU")
    p_shifts = nobet_map[p]

    ws3.cell(row=r, column=1, value=p).fill = fill_green
    ws3.cell(row=r, column=1).font = font_body
    ws3.cell(row=r, column=1).border = border

    ws3.cell(row=r, column=2, value="Mikro").fill = fill_green
    ws3.cell(row=r, column=2).font = font_body
    ws3.cell(row=r, column=2).alignment = align_center
    ws3.cell(row=r, column=2).border = border

    toplam_calisma = 0
    norm_gece = 0
    norm_normal = 0
    risk_gece = 0
    risk_normal = 0

    for d in range(1, 32):
        col_idx = d + 2
        cell = ws3.cell(row=r, column=col_idx)
        cell.border = border
        cell.alignment = align_center

        dt = datetime.date(2026, 10, d)
        is_off = dt.weekday() in [5, 6] or d == 29

        if is_off: cell.fill = fill_grey

        if is_muaf:
            if is_off: cell.value = "T"
            elif d == 28: cell.value = 5; toplam_calisma += 5
            else: cell.value = 8; toplam_calisma += 8
        else:
            if d in p_shifts:
                cell.value = 24
                cell.font = font_24
                if p_shifts[d] == "24A": cell.fill = fill_yellow

                n_saat = calculate_shift_hours_dynamic(d)
                toplam_calisma += n_saat

                # Gece/Gündüz Saat Hesabı
                g_saat = 12 if (is_off or d in [27, 28] or dt.weekday() == 4) else 8
                gunduz_saat = max(0, n_saat - g_saat)

                if p_shifts[d] == "24A":
                    risk_gece += g_saat
                    risk_normal += gunduz_saat
                else:
                    norm_gece += g_saat
                    norm_normal += gunduz_saat

            elif (d - 1) in p_shifts:
                if is_off: cell.value = "T"
                else: cell.value = "Nİ"; cell.font = font_ni; toplam_calisma += 8
            else:
                if is_off: cell.value = "T"
                elif d == 28: cell.value = 5; toplam_calisma += 5
                else: cell.value = 8; toplam_calisma += 8

    aylik_hedef = 165
    fazla_saat = max(0, toplam_calisma - aylik_hedef)

    metrics = [toplam_calisma, aylik_hedef, fazla_saat, norm_gece, norm_normal, risk_gece, risk_normal]
    for m_i, m_val in enumerate(metrics, 34):
        mc = ws3.cell(row=r, column=m_i, value=m_val)
        mc.font = font_bold
        mc.alignment = align_center
        mc.border = border

ws3.column_dimensions["A"].width = 28
ws3.column_dimensions["B"].width = 12
for d in range(1, 32):
    ws3.column_dimensions[get_column_letter(d + 2)].width = 4.2
for col_idx in range(34, 41):
    ws3.column_dimensions[get_column_letter(col_idx)].width = 6.5

output_filename = "Ekim_2026_Mevcut_Nobet_Listesi.xlsx"
wb.save(output_filename)
print(f"🎉 3 Sekmeli Tam Excel Dosyası Oluşturuldu: {output_filename}")
